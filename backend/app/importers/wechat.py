"""WeChat Pay statement (微信支付账单流水文件 .xlsx) importer.

Layout facts this parser relies on -- and deliberately does not:

* The export opens with a free-text preamble (nickname, period, notes) whose
  length varies. The header row is found by scanning for the required column
  names, never by a fixed row number.
* 金额(元) may arrive as a number cell or as text such as "¥190.93". It is read
  through Decimal (from the shortest float repr, never float arithmetic) and
  must have at most two decimals.
* 交易单号 is an identifier of up to ~32 digits. It must be a text cell: a
  numeric cell has already lost precision in Excel, so it is rejected rather
  than silently corrupted.
* 交易时间 is China time (Asia/Shanghai, UTC+08:00, no DST since 1991), whatever
  the timezone of the computer doing the import.
* "/" means "empty" in every column.

The content is parsed from memory; nothing is written to disk.
"""

from __future__ import annotations

import re
import unicodedata
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO

from app.importers.base import (
    Movement,
    ParsedStatement,
    RowError,
    StatementRow,
    UnsupportedStatementError,
    source_calendar_date,
)

#: WeChat statements are written in China Standard Time. UTC+08:00 with no DST
#: since 1991, so a fixed offset is exact and needs no tz database on Windows.
SOURCE_TIMEZONE = "Asia/Shanghai"
CHINA_TZ = timezone(timedelta(hours=8), SOURCE_TIMEZONE)
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
HEADER_SCAN_ROWS = 200

COL_TIME = "交易时间"
COL_TYPE = "交易类型"
COL_COUNTERPARTY = "交易对方"
COL_PRODUCT = "商品"
COL_DIRECTION = "收/支"
COL_AMOUNT = "金额(元)"
COL_METHOD = "支付方式"
COL_STATUS = "当前状态"
COL_ID = "交易单号"
COL_MERCHANT_ID = "商户单号"
COL_REMARK = "备注"

REQUIRED_COLUMNS = (
    COL_TIME, COL_TYPE, COL_COUNTERPARTY, COL_DIRECTION, COL_AMOUNT, COL_METHOD, COL_STATUS, COL_ID,
)
OPTIONAL_COLUMNS = (COL_PRODUCT, COL_MERCHANT_ID, COL_REMARK)

WALLET = "零钱"
WALLET_FUND = "零钱通"
WEALTH = "理财通"

#: Statuses of rows that moved no money (or whose money came straight back).
NOT_COMPLETED = ("支付失败", "已关闭", "已撤销", "已退还", "对方已退还", "朋友已退还", "已拒收")
TRANSFER_NOTE_PREFIX = re.compile(r"^(转账备注|付款方留言|收款方备注)[:：]\s*")
ID_PATTERN = re.compile(r"^[A-Za-z0-9_\-]{6,64}$")
TIME_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M")


def normalise_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"\s+", "", text)


def clean(value: object) -> str:
    """Cell to trimmed text; WeChat's "/" placeholder becomes empty."""
    if value is None:
        return ""
    text = str(value).replace("\t", " ").strip()
    return "" if text in ("/", "-", "--") else text


def parse_amount(value: object) -> int:
    """Amount cell -> fen, exactly. Raises ValueError."""
    if isinstance(value, bool) or value is None:
        raise ValueError("amount is empty")
    if isinstance(value, (int, float, Decimal)):
        text = repr(value) if isinstance(value, float) else str(value)
    else:
        text = unicodedata.normalize("NFKC", str(value)).strip()
        text = text.replace("¥", "").replace(",", "").replace(" ", "")
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"amount '{value}' is not a number") from exc
    if not amount.is_finite() or amount <= 0:
        raise ValueError("amount must be greater than zero")
    if amount.as_tuple().exponent < -2 and amount != amount.quantize(Decimal("0.01")):
        raise ValueError("amount has more than two decimal places")
    return int((amount * 100).to_integral_value())


def parse_time(value: object) -> datetime:
    """Transaction time -> aware datetime in China time. Raises ValueError."""
    if isinstance(value, datetime):
        if value.tzinfo is None:
            return value.replace(tzinfo=CHINA_TZ)
        return value.astimezone(CHINA_TZ)
    text = clean(value)
    for fmt in TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=CHINA_TZ)
        except ValueError:
            continue
    raise ValueError(f"transaction time '{text}' is not a date and time")


def parse_id(value: object) -> str:
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        raise ValueError(
            "交易单号 is stored as a number, so Excel has already rounded it. "
            "Re-export the statement without opening and re-saving it in Excel"
        )
    text = clean(value).strip("'\"`")
    if not ID_PATTERN.match(text):
        raise ValueError("交易单号 is missing or malformed")
    return text


def interpret(row: StatementRow) -> None:
    """Translate WeChat vocabulary into a movement plus funding labels.

    Only what the statement proves is concluded. Anything that could be income,
    spending or a move between the user's own money -- and cannot be told apart
    from the row alone -- is left for review, never guessed.
    """
    kind, status, method = row.source_type, row.status, row.payment_method
    funding = method or (status[3:] if status.startswith("已存入") else "")
    row.source_label = funding

    if any(marker in status for marker in NOT_COMPLETED):
        row.movement, row.review_reason = Movement.IGNORED, f"Not a completed payment ({status})"
        return
    if "退款" in kind or "退款" in status:
        row.movement = Movement.REVIEW
        row.review_reason = "Refund: decide whether to record it or skip it"
        return

    transfer = _own_funds_transfer(row)
    if transfer is not None:
        row.movement = Movement.TRANSFER
        row.source_label, row.dest_label = transfer
        return
    if row.direction == "income":
        row.movement = Movement.INCOME
    elif row.direction == "expense":
        row.movement = Movement.EXPENSE
    else:
        row.movement = Movement.REVIEW
        row.review_reason = f"Neutral WeChat transaction ({kind}): choose how to record it"


def _own_funds_transfer(row: StatementRow) -> tuple[str, str] | None:
    """(from, to) funding labels when the row moves money between the user's
    own WeChat balances and cards; None otherwise."""
    kind, method, other = row.source_type, row.payment_method, row.counterparty
    if match := re.match(r"^转入零钱通[-－]来自(.+)$", kind):
        return match.group(1).strip(), WALLET_FUND
    if match := re.match(r"^零钱通转出[-－]到(.+)$", kind):
        return WALLET_FUND, match.group(1).strip()
    if kind == "零钱充值":
        return method, WALLET
    if kind == "零钱提现":
        return WALLET, other or row.product
    if kind in ("零钱通转入", "转入零钱通"):
        return method or WALLET, WALLET_FUND
    if kind == "零钱通转出":
        return WALLET_FUND, method or WALLET
    if kind == "信用卡还款":
        return method, other or row.product
    if WEALTH in kind:
        redeem = any(word in kind for word in ("赎回", "卖出", "取出", "转出"))
        return (WEALTH, method or WALLET) if redeem else (method or WALLET, WEALTH)
    if row.direction == "neutral" and method and other:
        return method, other
    return None


def _cell_reader(cells: tuple, columns: dict[str, int]):
    def get(name: str) -> object:
        index = columns.get(name)
        return cells[index] if index is not None and index < len(cells) else None

    return get


class WeChatStatementImporter:
    source = "wechat"
    display_name = "WeChat Pay statement"
    currency = "CNY"
    timezone_name = SOURCE_TIMEZONE

    def parse(self, content: bytes) -> ParsedStatement:
        if len(content) > MAX_BYTES:
            raise UnsupportedStatementError("The file is larger than 10 MB.")
        if not content.startswith(b"PK\x03\x04"):
            raise UnsupportedStatementError(
                "This is not an .xlsx workbook. Export the WeChat Pay statement as Excel (.xlsx)."
            )
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # corrupt or not a spreadsheet
            raise UnsupportedStatementError("The workbook could not be opened.") from exc
        try:
            for sheet in workbook.worksheets:
                parsed = self._parse_sheet(sheet)
                if parsed is not None:
                    return parsed
        finally:
            workbook.close()
        raise UnsupportedStatementError(
            "Unsupported format: no sheet has the WeChat Pay columns "
            + "、".join(REQUIRED_COLUMNS)
            + "."
        )

    def _parse_sheet(self, sheet) -> ParsedStatement | None:
        rows = sheet.iter_rows(values_only=True)
        columns: dict[str, int] | None = None
        best_missing: list[str] = []
        header_row = 0
        for index, cells in enumerate(rows, start=1):
            names = [normalise_header(cell) for cell in cells]
            present = [name for name in REQUIRED_COLUMNS if name in names]
            if len(present) == len(REQUIRED_COLUMNS):
                wanted = (*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS)
                columns = {name: names.index(name) for name in wanted if name in names}
                header_row = index
                break
            if len(present) >= 3 and not best_missing:
                best_missing = [name for name in REQUIRED_COLUMNS if name not in names]
            if index >= HEADER_SCAN_ROWS:
                break
        if columns is None:
            if best_missing:
                raise UnsupportedStatementError(
                    "Unsupported format: the statement is missing required columns "
                    + "、".join(best_missing)
                    + "."
                )
            return None

        parsed = ParsedStatement(source=self.source, rows=[], header_row=header_row)
        for offset, cells in enumerate(rows, start=header_row + 1):
            if len(parsed.rows) + len(parsed.errors) >= MAX_ROWS:
                raise UnsupportedStatementError(f"The statement has more than {MAX_ROWS} rows.")
            if not any(clean(cell) for cell in cells):
                continue
            get = _cell_reader(cells, columns)
            if not clean(get(COL_ID)) and not any(
                clean(get(name)) for name in (COL_TYPE, COL_DIRECTION, COL_AMOUNT)
            ):
                continue  # trailing notes or separators, not a transaction
            try:
                parsed.rows.append(self._row(offset, get))
            except ValueError as exc:
                parsed.errors.append(RowError(offset, str(exc)))
        return parsed

    def _row(self, row_number: int, get) -> StatementRow:
        occurred_at = parse_time(get(COL_TIME))
        raw_direction = clean(get(COL_DIRECTION))
        direction = {"收入": "income", "支出": "expense"}.get(raw_direction, "neutral")
        product = clean(get(COL_PRODUCT))
        note = ""
        if TRANSFER_NOTE_PREFIX.match(product):
            note = TRANSFER_NOTE_PREFIX.sub("", product)
            product = ""
        row = StatementRow(
            row_number=row_number,
            external_id=parse_id(get(COL_ID)),
            occurred_at=occurred_at,
            source_date=source_calendar_date(occurred_at, CHINA_TZ),
            source_type=clean(get(COL_TYPE)),
            direction=direction,
            counterparty=clean(get(COL_COUNTERPARTY)),
            product=product,
            note=note,
            remark=clean(get(COL_REMARK)),
            amount_minor=parse_amount(get(COL_AMOUNT)),
            currency=self.currency,
            payment_method=clean(get(COL_METHOD)),
            status=clean(get(COL_STATUS)),
            merchant_order_id=clean(get(COL_MERCHANT_ID)),
            raw_direction=raw_direction,
            source_timezone=SOURCE_TIMEZONE,
        )
        interpret(row)
        return row
